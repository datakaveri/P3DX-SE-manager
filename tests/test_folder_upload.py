"""Folder upload: one tar through the unchanged transport, two modes.

Joint (csv/json/excel): every member combined into one input, SKALD once.
Per file (dicom/image): each member through the single-file path, outputs in a
tar with _manifest.json.

The tests that need SKALD itself run the real pipeline binary when one is
available — set SKALD_PIPELINE_BIN, or build ~/k-anonymisation/SKALD — and
skip otherwise. The rest need nothing but the standard library (openpyxl for
the Excel cases).

Run:  python3 -m unittest discover -s tests
      P3DX_SCRATCH_FULL=1 ... to measure scratch at the full 300 MB
"""

from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import os
import shutil
import subprocess
import sys
import tarfile
import tempfile
import types
import unittest
from collections import Counter
from unittest import mock

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)
sys.path.insert(0, os.path.join(_ROOT, "Fetch_data"))

if "dotenv" not in sys.modules:
    _stub = types.ModuleType("dotenv")
    _stub.load_dotenv = lambda *a, **k: None
    sys.modules["dotenv"] = _stub

from cryptography.hazmat.primitives import hashes  # noqa: E402
from cryptography.hazmat.primitives.asymmetric import padding, rsa  # noqa: E402

from lib import direct_upload, folder_bundle  # noqa: E402
from lib.config import config  # noqa: E402
from lib.direct_upload import head_matches_format  # noqa: E402
from lib.folder_bundle import FolderError, MemberFailure  # noqa: E402
from enclave.enclave_direct_upload import (  # noqa: E402
    MAX_FOLDER_BYTES, MAX_FOLDER_FILES, MB, UploadError, UploadManager,
    chunk_digest_root, encrypt_chunk,
)

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

SKALD_BIN = os.environ.get("SKALD_PIPELINE_BIN") or os.path.expanduser(
    "~/k-anonymisation/SKALD/target/release/skald_pipeline")
HAVE_SKALD = os.access(SKALD_BIN, os.X_OK)

DICOM = b"\x00" * 128 + b"DICM" + b"\x02\x00\x10\x00UI" + b"\x00" * 64
PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
XLSX_MAGIC = b"PK\x03\x04" + b"\x00" * 64


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #


def build_tar(path, members, fmt=tarfile.PAX_FORMAT):
    """members: (name, bytes) or a callable(tar) for a hand-built entry."""
    with tarfile.open(path, "w:", format=fmt) as tf:
        for m in members:
            if callable(m):
                m(tf)
                continue
            name, data = m
            info = tarfile.TarInfo(name)
            info.size, info.mode = len(data), 0o644
            tf.addfile(info, io.BytesIO(data))
    return path


def _read(path):
    with open(path, "rb") as f:
        return f.read()


def skald_section(**over):
    section = {
        "output_path": "generalized.csv",
        "output_directory": "output",
        "output_format": "match_input",
        "suppression_limit": 1.0,
    }
    section.update(over)
    return section


def run_skald(workdir, data_name, data_path, cfg):
    """Run the real SKALD binary over one input, as its container does."""
    for d in ("config", "data", "output"):
        os.makedirs(os.path.join(workdir, d), exist_ok=True)
    shutil.copy(data_path, os.path.join(workdir, "data", data_name))
    with open(os.path.join(workdir, "config", "config.json"), "w") as f:
        json.dump(cfg, f)
    proc = subprocess.run([SKALD_BIN], cwd=workdir, capture_output=True, text=True, timeout=300)
    with open(os.path.join(workdir, "output", "status.json")) as f:
        status = json.load(f)
    if proc.returncode != 0 or status.get("status") != "success":
        raise AssertionError(f"SKALD failed: {status}\n{proc.stdout[-2000:]}")
    return status


def read_output_rows(workdir, ext):
    path = os.path.join(workdir, "output", f"generalized{ext}")
    if ext == ".json":
        with open(path) as f:
            return json.load(f)
    with open(path, newline="") as f:
        return list(csv.DictReader(f))


def classes(rows, qis):
    return Counter(tuple(r[q] for q in qis) for r in rows)


class TempDirCase(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="p3dx-folder-")
        self.addCleanup(shutil.rmtree, self.dir, ignore_errors=True)

    def path(self, *parts):
        p = os.path.join(self.dir, *parts)
        os.makedirs(os.path.dirname(p), exist_ok=True)
        return p

    def validate(self, members, fmt="csv", tar_fmt=tarfile.PAX_FORMAT):
        tar = build_tar(self.path("in.tar"), members, tar_fmt)
        return folder_bundle.validate_archive(tar, fmt, head_matches_format)


# --------------------------------------------------------------------------- #
# 1-3. detection, the /upload/init cap, archive validation
# --------------------------------------------------------------------------- #


class Detection(TempDirCase):
    def test_tar_name_and_magic_is_a_folder(self):
        tar = build_tar(self.path("scans.tar"), [("IM0001", DICOM)])
        self.assertTrue(folder_bundle.detect("scans.tar", tar, "dicom"))
        self.assertTrue(folder_bundle.detect("SCANS.TAR", tar, "dicom"))

    def test_a_tar_name_without_magic_fails(self):
        p = self.path("fake.tar")
        with open(p, "wb") as f:
            f.write(b"id,age\n1,2\n" * 100)
        with self.assertRaisesRegex(FolderError, "not a tar archive"):
            folder_bundle.detect("fake.tar", p, "csv")

    def test_magic_without_a_tar_name_is_a_single_file(self):
        tar = build_tar(self.path("x.csv"), [("a.csv", b"id\n1\n")])
        self.assertFalse(folder_bundle.detect("x.csv", tar, "csv"))

    def test_format_tar_fails(self):
        tar = build_tar(self.path("x.tar"), [("a.csv", b"id\n1\n")])
        with self.assertRaisesRegex(FolderError, "not a folder member format"):
            folder_bundle.detect("x.tar", tar, "tar")


class _Uploads:
    """A real UploadManager, driven the way the middleware drives it."""

    def __init__(self, scratch):
        self.key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
        self.manager = UploadManager(self.key, scratch_dir=scratch)

    def wrap(self, raw):
        return base64.b64encode(self.key.public_key().encrypt(raw, padding.OAEP(
            mgf=padding.MGF1(hashes.SHA256()), algorithm=hashes.SHA256(), label=None))).decode()

    def init_body(self, filename, fmt, total, chunk=None):
        chunk = chunk or max(1, total)
        return {
            "filename": filename, "format": fmt, "total_bytes": total,
            "chunk_size": min(chunk, 64 * MB), "total_chunks": -(-total // min(chunk, 64 * MB)),
            "cipher": "AES-256-GCM", "key_wrap": "RSA-OAEP-SHA256",
            "wrapped_key": self.wrap(os.urandom(32)), "base_iv": base64.b64encode(os.urandom(12)).decode(),
            "output_wrapped_key": self.wrap(os.urandom(32)),
            "output_base_iv": base64.b64encode(os.urandom(12)).decode(),
        }

    def upload(self, user, filename, fmt, payload):
        dkey, iv = os.urandom(32), os.urandom(12)
        body = self.init_body(filename, fmt, len(payload))
        body["wrapped_key"] = self.wrap(dkey)
        body["base_iv"] = base64.b64encode(iv).decode()
        uid = self.manager.init(user, body)["upload_id"]
        digest = hashlib.sha256(payload).hexdigest()
        self.manager.put_chunk(user, uid, 0, encrypt_chunk(dkey, iv, uid, 0, 1, payload), digest)
        self.manager.complete(user, uid, chunk_digest_root([digest]))
        return f"enclave://upload/{uid}"


class InitCap(TempDirCase):
    def setUp(self):
        super().setUp()
        self.up = _Uploads(os.path.join(self.dir, "scratch"))

    def init(self, filename, fmt, total):
        return self.up.manager.init(f"user-{os.urandom(4).hex()}", self.up.init_body(
            filename, fmt, total, chunk=64 * MB))

    def test_an_image_folder_is_held_to_the_folder_cap_not_the_image_cap(self):
        self.init("photos.tar", "image", 200 * MB)          # 8x the 25 MB image cap

    def test_the_folder_cap_is_300_mb_whatever_the_format(self):
        for fmt in ("csv", "excel", "image", "dicom"):
            with self.subTest(fmt=fmt):
                self.init(f"f-{fmt}.tar", fmt, MAX_FOLDER_BYTES)
                with self.assertRaises(UploadError) as ctx:
                    self.init(f"g-{fmt}.tar", fmt, MAX_FOLDER_BYTES + 1)
                self.assertEqual(ctx.exception.status, 413)

    def test_single_files_keep_their_per_format_cap(self):
        with self.assertRaises(UploadError) as ctx:
            self.init("big.png", "image", 26 * MB)
        self.assertEqual(ctx.exception.status, 413)

    def test_format_tar_is_rejected_at_init(self):
        with self.assertRaises(UploadError) as ctx:
            self.init("x.tar", "tar", 1024)
        self.assertEqual(ctx.exception.status, 400)

    def test_constants_match_the_ui(self):
        self.assertEqual(MAX_FOLDER_BYTES, 300 * MB)
        self.assertEqual(MAX_FOLDER_FILES, 10_000)


class Validation(TempDirCase):
    def test_a_well_formed_folder_passes(self):
        long_name = "sub/" + "é" * 80 + ".csv"          # pax path: > 100 bytes, non-ASCII
        s = self.validate([("a.csv", b"id\n1\n"), (long_name, b"id\n2\n")])
        self.assertEqual(s.files_total, 2)
        self.assertEqual(s.mode, "joint")

    def fails(self, members, pattern, fmt="csv", tar_fmt=tarfile.PAX_FORMAT):
        with self.assertRaisesRegex(FolderError, pattern):
            self.validate(members, fmt, tar_fmt)

    def special(self, name, type_, **kw):
        def add(tf):
            info = tarfile.TarInfo(name)
            info.type = type_
            for k, v in kw.items():
                setattr(info, k, v)
            tf.addfile(info)
        return add

    def test_symlink(self):
        self.fails([("a.csv", b"id\n1\n"), self.special("b.csv", tarfile.SYMTYPE, linkname="/etc/passwd")],
                   "not a regular file")

    def test_hardlink(self):
        self.fails([("a.csv", b"id\n1\n"), self.special("b.csv", tarfile.LNKTYPE, linkname="a.csv")],
                   "not a regular file")

    def test_directory(self):
        self.fails([self.special("sub", tarfile.DIRTYPE)], "not a regular file")

    def test_device(self):
        self.fails([self.special("dev.csv", tarfile.CHRTYPE, devmajor=1, devminor=3)], "not a regular file")

    def test_fifo(self):
        self.fails([self.special("p.csv", tarfile.FIFOTYPE)], "not a regular file")

    def test_sparse(self):
        self.fails([self.special("s.csv", tarfile.GNUTYPE_SPARSE)], "not a regular file",
                   tar_fmt=tarfile.GNU_FORMAT)

    def test_gnu_long_name(self):
        self.fails([("d/" + "x" * 120 + ".csv", b"id\n1\n")], "unsupported tar header",
                   tar_fmt=tarfile.GNU_FORMAT)

    def test_dot_dot(self):
        self.fails([("../x.csv", b"id\n1\n")], r"'\.\.'")

    def test_absolute_path(self):
        self.fails([("/etc/x.csv", b"id\n1\n")], "absolute path")

    def test_nul_in_path(self):
        def add(tf):
            info = tarfile.TarInfo("a.csv")
            info.size = 5
            info.pax_headers = {"path": "a\x00b.csv"}
            tf.addfile(info, io.BytesIO(b"id\n1\n"))
        self.fails([add], "NUL")

    def test_duplicate_path(self):
        self.fails([("a.csv", b"id\n1\n"), ("a.csv", b"id\n2\n")], "appears twice")

    def test_member_over_its_cap(self):
        with mock.patch.dict(folder_bundle.MAX_TOTAL_BYTES_BY_FORMAT, {"csv": 10}):
            self.fails([("a.csv", b"id\n1\n"), ("b.csv", b"id\n" + b"1\n" * 10)], "over the")

    def test_sum_over_the_folder_cap(self):
        with mock.patch.object(folder_bundle, "MAX_FOLDER_BYTES", 20):
            self.fails([("a.csv", b"id\n" + b"1\n" * 5), ("b.csv", b"id\n" + b"1\n" * 5)],
                       "add up to more than")

    def test_10001_members(self):
        members = [(f"m{i:05d}.csv", b"i\n1\n") for i in range(MAX_FOLDER_FILES + 1)]
        self.fails(members, "more than 10,000 files")

    def test_exactly_10000_members_pass(self):
        members = [(f"m{i:05d}.csv", b"i\n1\n") for i in range(MAX_FOLDER_FILES)]
        self.assertEqual(self.validate(members).files_total, MAX_FOLDER_FILES)

    def test_empty_archive(self):
        self.fails([], "empty|not a readable")

    def test_a_csv_member_that_is_really_a_zip(self):
        self.fails([("a.csv", b"id\n1\n"), ("b.csv", XLSX_MAGIC)], "does not match the declared format")

    def test_every_member_is_sniffed_not_just_the_first(self):
        self.fails([("IM1", DICOM), ("IM2", DICOM), ("IM3", PNG)], "IM3", fmt="dicom")

    def test_dicom_members_without_an_extension_pass(self):
        self.assertEqual(self.validate([("s1/IM0001", DICOM), ("s1/IM0002", DICOM)], "dicom").mode,
                         "per_file")

    def test_trailing_garbage(self):
        tar = build_tar(self.path("in.tar"), [("a.csv", b"id\n1\n")])
        with open(tar, "ab") as f:
            f.write(b"\x00" * 512 + b"hidden" + b"\x00" * 506)
        with self.assertRaisesRegex(FolderError, "data after its last file|not a readable"):
            folder_bundle.validate_archive(tar, "csv", head_matches_format)


class Staging(TempDirCase):
    """stage_for_pipeline: the bundle-upload check, over a real upload session."""

    def setUp(self):
        super().setUp()
        self.scratch = os.path.join(self.dir, "scratch")
        self.up = _Uploads(self.scratch)
        p = mock.patch.object(direct_upload, "get_manager", return_value=self.up.manager)
        p.start()
        self.addCleanup(p.stop)

    def stage(self, filename, fmt, payload, application=None):
        ref = self.up.upload("u1", filename, fmt, payload)
        direct_upload.stage_for_pipeline("u1", ref, application=application)
        uid = ref.rsplit("/", 1)[1]
        return direct_upload.read_staged_meta(self.scratch, uid)

    def test_a_folder_is_recorded_in_the_sidecar(self):
        tar = _read(build_tar(self.path("t.tar"), [("a.json", b'{"id":1}'), ("b.json", b'{"id":2}')]))
        meta = self.stage("patients.tar", "json", tar)
        self.assertEqual(meta["folder"], {"mode": "joint", "files_total": 2})
        self.assertEqual(meta["format"], "json")

    def test_a_bad_archive_rejects_the_bundle(self):
        tar = _read(build_tar(self.path("t.tar"), [("../a.csv", b"id\n1\n")]))
        with self.assertRaises(UploadError) as ctx:
            self.stage("x.tar", "csv", tar)
        self.assertEqual(ctx.exception.status, 422)

    def test_dp_with_a_folder_fails_with_the_dp_message(self):
        tar = _read(build_tar(self.path("t.tar"), [("a.csv", b"id\n1\n")]))
        with self.assertRaises(UploadError) as ctx:
            self.stage("x.tar", "csv", tar, application="dp")
        self.assertEqual(ctx.exception.description,
                         "folder upload is not supported for differential privacy yet")

    def test_dp_single_file_is_unaffected(self):
        meta = self.stage("x.csv", "csv", b"id,age\n1,2\n", application="dp")
        self.assertNotIn("folder", meta)

    def test_single_file_sidecar_is_unchanged(self):
        meta = self.stage("scan.csv", "csv", b"id,age\n1,34\n")
        self.assertEqual(meta, {"filename": "scan.csv", "format": "csv"})

    def test_single_file_bytes_are_staged_unchanged(self):
        payload = b"id,age\n1,34\n2,41\n"
        ref = self.up.upload("u1", "scan.csv", "csv", payload)
        direct_upload.stage_for_pipeline("u1", ref)
        data_dir = os.path.join(self.dir, "data")
        import fetch_data as fd  # noqa: PLC0415 — sys.path set by direct_upload
        with mock.patch.object(direct_upload, "_scratch_dir", return_value=self.scratch):
            fd._stage_direct_upload_input(ref, data_dir, "csv")
        with open(os.path.join(data_dir, "scan.csv"), "rb") as f:
            self.assertEqual(f.read(), payload)


class ApplicationDetection(unittest.TestCase):
    def test_compose_urls(self):
        import P3DX_SDK
        base = "https://raw.githubusercontent.com/datakaveri/Docker-Compose/main"
        for d, app in (("differential-privacy", "dp"), ("skald", "skald"),
                       ("skald-dicom", "skald_dicom"), ("skald-image", "skald_image")):
            self.assertEqual(P3DX_SDK.application_from_compose_url(f"{base}/{d}/docker-compose.yml"), app)
        self.assertIsNone(P3DX_SDK.application_from_compose_url(None))


# --------------------------------------------------------------------------- #
# 4. Joint
# --------------------------------------------------------------------------- #


class JointBase(TempDirCase):
    def combine(self, members, fmt, section, data_type="patients"):
        tar = build_tar(self.path("in.tar"), members)
        ext = {"csv": ".csv", "json": ".json", "excel": ".xlsx"}[fmt]
        out = self.path("data", data_type + ext)
        result = folder_bundle.combine_tabular(tar, fmt, section, out, os.path.dirname(out))
        self.assertFalse(os.path.exists(os.path.join(os.path.dirname(out), ".folder-records.spool")),
                         "the record spool must not outlive the combine")
        return result


def _patient(i):
    rec = {"patient_id": f"P{i:03d}", "age": 20 + 2 * i, "gender": "MF"[i % 2],
           "diagnosis": ["flu", "cold", "asthma"][i % 3]}
    if i % 3 == 0:
        rec["allergy"] = "penicillin"                   # optional field
    if i % 4 == 0:
        rec["visits"] = [{"date": "2024-01-0%d" % (1 + i % 9), "ward": "A"}]  # array field
    return rec


class JointJson(JointBase):
    SECTION = skald_section(
        k_anonymize={"k": 5},
        hashing_with_salt=["patient_id"],
        quasi_identifiers={"numerical": [{"column": "age", "type": "int"}], "categorical": ["gender"]},
        size={"age": 2},
        sensitive_parameter="diagnosis",
        insensitive_columns=["allergy", "visits"],
    )

    def members(self):
        return [(f"patients/p{i:03d}.json", json.dumps(_patient(i)).encode()) for i in range(20)]

    def test_one_record_per_object_file_and_nothing_from_the_file_names(self):
        result = self.combine(self.members(), "json", self.SECTION)
        with open(result.path, "rb") as f:
            raw = f.read()
        records = json.loads(raw)
        self.assertEqual(result.records_in, 20)
        self.assertEqual(result.files_total, 20)
        self.assertEqual(len(records), 20, "an object with an array field is ONE record")
        self.assertNotIn(b"patients/", raw)
        self.assertNotIn(b".json", raw)
        for rec in records:
            self.assertLessEqual(set(rec), {"patient_id", "age", "gender", "diagnosis", "allergy", "visits"})
        order = [r["patient_id"] for r in records]
        self.assertNotEqual(order, sorted(order), "rows must not keep file-path order")
        self.assertEqual(sorted(order), [f"P{i:03d}" for i in range(20)])

    def test_a_top_level_array_is_a_list_of_records(self):
        members = [("a.json", json.dumps([_patient(0), _patient(1)]).encode()),
                   ("b.json", json.dumps(_patient(2)).encode())]
        self.assertEqual(self.combine(members, "json", self.SECTION).records_in, 3)

    @unittest.skipUnless(HAVE_SKALD, "SKALD pipeline binary not available")
    def test_k5_holds_over_the_whole_folder(self):
        result = self.combine(self.members(), "json", self.SECTION)
        work = os.path.join(self.dir, "skald")
        run_skald(work, "patients.json", result.path, {"data_type": "patients", "patients": self.SECTION})
        rows = read_output_rows(work, ".json")
        self.assertLessEqual(len(rows), 20)
        for key, n in classes(rows, ("age", "gender")).items():
            self.assertGreaterEqual(n, 5, f"class {key} has {n} < k records")
        blob = json.dumps(rows)
        self.assertNotIn("p000", blob)
        self.assertNotIn("source", blob)
        self.assertTrue(all(len(r["patient_id"]) == 64 for r in rows), "IDs are salted hashes")


class JointCsv(JointBase):
    # Categorical QI, no hierarchy: the only generalisation is to "*", so the
    # count of starred rows is exactly what k cost.
    SECTION = skald_section(
        k_anonymize={"k": 3},
        flow_mode="direct",          # SKALD's OLA-1 ladders cover only a few named columns
        quasi_identifiers={"categorical": ["zone"]},
        sensitive_parameter="diagnosis",
        insensitive_columns=["id"],
    )

    def files(self):
        # Each file holds each zone ONCE: no file reaches k=3 on its own; the
        # three together hold every zone three times.
        return [(f"site{s}.csv", "".join(
            ["id,zone,diagnosis\n"] + [f"{s}{z},{z},d{z}\n" for z in "ABCD"]).encode())
            for s in range(3)]

    def test_combines_over_all_files(self):
        result = self.combine(self.files(), "csv", self.SECTION)
        with open(result.path, newline="") as f:
            rows = list(csv.DictReader(f))
        self.assertEqual(len(rows), 12)
        self.assertEqual(Counter(r["zone"] for r in rows), {"A": 3, "B": 3, "C": 3, "D": 3})

    def test_union_of_columns_fills_missing_cells_empty(self):
        section = dict(self.SECTION, insensitive_columns=["id", "note"])
        members = self.files() + [("extra.csv", b"id,zone,diagnosis,note\n9,A,dA,x\n")]
        result = self.combine(members, "csv", section)
        with open(result.path, newline="") as f:
            rows = list(csv.DictReader(f))
        self.assertEqual([r for r in rows if r["id"] == "9"][0]["note"], "x")
        self.assertTrue(all(r["note"] == "" for r in rows if r["id"] != "9"))

    @unittest.skipUnless(HAVE_SKALD, "SKALD pipeline binary not available")
    def test_joint_reaches_k_and_suppresses_less_than_per_file(self):
        cfg = {"data_type": "sites", "sites": self.SECTION}
        result = self.combine(self.files(), "csv", self.SECTION, data_type="sites")
        joint = os.path.join(self.dir, "joint")
        run_skald(joint, "sites.csv", result.path, cfg)
        joint_rows = read_output_rows(joint, ".csv")
        joint_starred = sum(r["zone"] == "*" for r in joint_rows)

        per_file_starred = 0
        for i, (name, data) in enumerate(self.files()):
            src = self.path(f"pf{i}", name)
            with open(src, "wb") as f:
                f.write(data)
            work = os.path.join(self.dir, f"pf{i}-run")
            run_skald(work, name, src, cfg)
            per_file_starred += sum(r["zone"] == "*" for r in read_output_rows(work, ".csv"))

        self.assertEqual(joint_starred, 0)
        self.assertEqual(per_file_starred, 12)
        for key, n in classes(joint_rows, ("zone",)).items():
            self.assertGreaterEqual(n, 3)

    @unittest.skipUnless(HAVE_SKALD, "SKALD pipeline binary not available")
    def test_the_same_value_in_two_files_hashes_the_same(self):
        section = skald_section(k_anonymize={"k": 2}, hashing_with_salt=["patient_id"],
                                flow_mode="direct", quasi_identifiers={"categorical": ["zone"]},
                                insensitive_columns=["visit"])
        members = [
            ("a.csv", b"patient_id,zone,visit\nP001,A,1\nP002,A,1\nP003,B,1\n"),
            ("b.csv", b"patient_id,zone,visit\nP001,B,2\nP004,B,2\n"),
        ]
        result = self.combine(members, "csv", section)
        work = os.path.join(self.dir, "salt")
        run_skald(work, "patients.csv", result.path, {"data_type": "patients", "patients": section})
        rows = read_output_rows(work, ".csv")
        hashes = Counter(r["patient_id"] for r in rows)
        self.assertEqual(len(hashes), 4, "four distinct patients -> four distinct hashes")
        self.assertEqual(sorted(hashes.values()), [1, 1, 1, 2], "P001 from both files -> one hash")
        self.assertNotIn("P001", json.dumps(rows))


class JointFailures(JointBase):
    SECTION = skald_section(quasi_identifiers={"numerical": [{"column": "age", "type": "int"}]},
                            insensitive_columns=["id"])

    def test_a_column_missing_from_the_config_fails_and_is_listed(self):
        members = [("a.csv", b"id,age\n1,30\n"), ("b.csv", b"id,age,postcode,ssn\n2,40,560001,x\n")]
        with self.assertRaises(FolderError) as ctx:
            self.combine(members, "csv", self.SECTION)
        self.assertIn("['postcode', 'ssn']", str(ctx.exception))

    def test_an_unreadable_member_fails_the_whole_job_and_is_named(self):
        members = [("ok.json", b'{"id": 1, "age": 30}'), ("broken/p2.json", b'{"id": 2, "age": ')]
        with self.assertRaises(FolderError) as ctx:
            self.combine(members, "json", self.SECTION)
        self.assertIn("broken/p2.json", str(ctx.exception))

    def test_a_non_utf8_csv_member_fails_and_is_named(self):
        members = [("a.csv", b"id,age\n1,30\n"), ("latin.csv", b"id,age\n\xe9,30\n")]
        with self.assertRaisesRegex(FolderError, "latin.csv"):
            self.combine(members, "csv", self.SECTION)

    def test_a_value_that_cannot_be_coerced_names_the_member_not_the_value(self):
        members = [("a.csv", b"id,age\n1,30\n"), ("b.csv", b"id,age\n2,SECRETVALUE\n")]
        with self.assertRaises(FolderError) as ctx:
            self.combine(members, "csv", self.SECTION)
        self.assertIn("b.csv", str(ctx.exception))
        self.assertNotIn("SECRETVALUE", str(ctx.exception))

    def test_a_record_missing_a_quasi_identifier_is_refused(self):
        members = [("a.json", b'{"id": 1, "age": 30}'), ("b.json", b'{"id": 2}')]
        with self.assertRaisesRegex(FolderError, "'age' is missing from 1 file"):
            self.combine(members, "json", self.SECTION)

    def test_a_member_row_wider_than_its_header_fails(self):
        with self.assertRaisesRegex(FolderError, "b.csv"):
            self.combine([("a.csv", b"id,age\n1,30\n"), ("b.csv", b"id,age\n2,40,extra\n")],
                         "csv", self.SECTION)


@unittest.skipUnless(openpyxl, "openpyxl not installed")
class JointExcel(JointBase):
    SECTION = skald_section(quasi_identifiers={"categorical": ["gender"]},
                            sensitive_parameter="diagnosis", join_keys=["id"],
                            insensitive_columns=["id", "age"])

    def workbook(self, patients, visits):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "patients"
        ws.append(["id", "age", "gender"])
        for row in patients:
            ws.append(row)
        ws2 = wb.create_sheet("visits")
        ws2.append(["id", "diagnosis"])
        for row in visits:
            ws2.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def members(self):
        # Both workbooks reuse ids 1 and 2: stacking first and joining after
        # would pair rows ACROSS workbooks.
        return [
            ("clinic_a.xlsx", self.workbook([[1, 30, "F"], [2, 41, "M"]], [[1, "flu"], [2, "cold"]])),
            ("clinic_b.xlsx", self.workbook([[1, 55, "M"], [2, 62, "F"]], [[1, "asthma"], [2, "gout"]])),
        ]

    def rows(self, result):
        wb = openpyxl.load_workbook(result.path, read_only=True)
        self.assertEqual(len(wb.worksheets), 1, "one sheet")
        values = list(wb.worksheets[0].iter_rows(values_only=True))
        header, body = values[0], values[1:]
        return [dict(zip(header, r)) for r in body]

    def test_each_workbook_is_joined_then_the_results_concatenated(self):
        result = self.combine(self.members(), "excel", self.SECTION)
        got = {(r["id"], r["age"], r["gender"], r["diagnosis"]) for r in self.rows(result)}
        self.assertEqual(got, {(1, 30, "F", "flu"), (2, 41, "M", "cold"),
                               (1, 55, "M", "asthma"), (2, 62, "F", "gout")})
        self.assertEqual(result.records_in, 4)

    def test_explicit_sheet_joins_are_honoured(self):
        section = dict(self.SECTION, sheet_joins=[{"left": "patients", "right": "visits", "on": "id"}])
        result = self.combine(self.members(), "excel", section)
        self.assertEqual(len(self.rows(result)), 4)

    def test_a_workbook_with_an_extra_column_fails(self):
        extra = openpyxl.Workbook()
        extra.active.append(["id", "age", "gender", "ssn"])
        extra.active.append([3, 70, "F", "x"])
        buf = io.BytesIO()
        extra.save(buf)
        with self.assertRaisesRegex(FolderError, r"\['ssn'\]"):
            self.combine(self.members() + [("c.xlsx", buf.getvalue())], "excel", self.SECTION)

    @unittest.skipUnless(HAVE_SKALD, "SKALD pipeline binary not available")
    def test_skald_reads_the_combined_workbook(self):
        section = dict(self.SECTION, k_anonymize={"k": 2})
        result = self.combine(self.members(), "excel", section)
        work = os.path.join(self.dir, "xl")
        run_skald(work, "patients.xlsx", result.path, {"data_type": "patients", "patients": section})
        self.assertTrue(os.path.exists(os.path.join(work, "output", "generalized.xlsx")))


class JointStagingEndToEnd(JointBase):
    """fetch_data._stage_folder_input: config lookup under data_type, combine,
    release of the bundle, and the counts handed on to status."""

    def test_stage_combines_releases_and_records_counts(self):
        import fetch_data as fd  # noqa: PLC0415
        cfg_dir, data_dir, scratch = (os.path.join(self.dir, d) for d in ("cfg", "data", "scratch"))
        for d in (cfg_dir, data_dir, scratch):
            os.makedirs(d)
        section = JointJson.SECTION
        with open(os.path.join(cfg_dir, "generated-config.json"), "w") as f:
            json.dump({"data_type": "my_patients", "my_patients": section}, f)
        uid = "0f8fad5b-d9cb-469f-a165-70867728950e"
        tar = build_tar(os.path.join(scratch, f"{uid}.bin"),
                        [(f"p{i}.json", json.dumps(_patient(i)).encode()) for i in range(6)])
        meta = {"filename": "my.patients.tar", "format": "json",
                "folder": {"mode": "joint", "files_total": 6}}
        with open(os.path.join(scratch, f"{uid}.meta.json"), "w") as f:
            json.dump(meta, f)

        with mock.patch.object(config.paths, "tee_input_config", cfg_dir), \
                mock.patch.object(direct_upload, "_scratch_dir", return_value=scratch):
            fd._stage_folder_input(tar, scratch, uid, meta, "json", data_dir)

        self.assertFalse(os.path.exists(tar), "bundle released once every member is loaded")
        self.assertEqual(os.listdir(data_dir), ["my.patients.json"])
        stored = direct_upload.read_staged_meta(scratch, uid)
        self.assertEqual(stored["folder"], {"mode": "joint", "files_total": 6, "records_in": 6})

    def test_dp_is_refused_in_the_subprocess_too(self):
        import fetch_data as fd  # noqa: PLC0415
        import P3DX_SDK  # noqa: PLC0415
        meta = {"filename": "x.tar", "format": "csv", "folder": {"mode": "joint", "files_total": 1}}
        with mock.patch.object(P3DX_SDK, "_compose_url",
                               "https://x/Docker-Compose/main/differential-privacy/docker-compose.yml"):
            with self.assertRaisesRegex(ValueError, "not supported for differential privacy"):
                fd._stage_folder_input("/nonexistent", self.dir, "u", meta, "csv", self.dir)


# --------------------------------------------------------------------------- #
# 5. Per file
# --------------------------------------------------------------------------- #


def fake_dicom_app(workdir):
    """Stands in for SKALD-DICOM: fails an unparsable file, otherwise writes a
    de-identified file that shares no bytes with the input."""
    def process(name, fh):
        data = fh.read()
        if data[128:132] != b"DICM":
            raise MemberFailure("unparsable DICOM")
        out = os.path.join(workdir, "after_deidentification.dcm")
        with open(out, "wb") as f:
            f.write(b"\x00" * 128 + b"DICM" + hashlib.sha256(data).digest())
        return out

    def cleanup():
        for n in os.listdir(workdir):
            os.unlink(os.path.join(workdir, n))
    return process, cleanup


class PerFile(TempDirCase):
    RAW_MARKER = b"RAW-PHI-MARKER-" + os.urandom(8).hex().encode()

    def members(self):
        return [("s1/IM0001", DICOM + b"a"), ("s1/IM0002", DICOM + b"b"),
                ("s1/IM0003", b"not dicom at all " + self.RAW_MARKER * 20)]

    def run_folder(self):
        tar = build_tar(self.path("in.tar"), self.members())
        work = os.path.join(self.dir, "work")
        os.makedirs(work)
        out = self.path("out.tar")
        process, cleanup = fake_dicom_app(work)
        return folder_bundle.run_per_file(tar, "dicom", out, process, cleanup), out

    def test_output_mirrors_input_paths_and_carries_a_manifest(self):
        manifest, out = self.run_folder()
        with tarfile.open(out) as tf:
            names = tf.getnames()
            inside = json.load(tf.extractfile("_manifest.json"))
        self.assertEqual(sorted(names), ["_manifest.json", "s1/IM0001_anonymised.dcm",
                                         "s1/IM0002_anonymised.dcm"])
        self.assertEqual(inside, manifest)
        self.assertEqual(manifest["mode"], "per_file")
        self.assertEqual(manifest["application"], "skald_dicom")
        self.assertIs(manifest["paths_deidentified"], False)
        self.assertEqual((manifest["files_total"], manifest["files_succeeded"],
                          manifest["files_failed"]), (3, 2, 1))
        self.assertEqual(manifest["files"][2],
                         {"path": "s1/IM0003", "status": "failed", "error": "unparsable DICOM"})

    def test_a_failed_members_raw_bytes_appear_nowhere(self):
        _, out = self.run_folder()
        with open(out, "rb") as f:
            self.assertNotIn(self.RAW_MARKER, f.read())

    def test_an_unexpected_error_never_leaks_its_text(self):
        tar = build_tar(self.path("in.tar"), [("a", DICOM), ("b", DICOM)])

        def process(name, fh):
            if name == "b":
                raise ValueError("tag (0010,0010) PatientName=JOHN^DOE could not be parsed")
            p = self.path("o.dcm")
            with open(p, "wb") as f:
                f.write(b"ok")
            return p
        manifest = folder_bundle.run_per_file(tar, "dicom", self.path("out.tar"), process, lambda: None)
        self.assertEqual(manifest["files"][1]["error"], "processing failed")

    def test_no_member_succeeding_fails_the_job(self):
        tar = build_tar(self.path("in.tar"), [("a", b"junk")])

        def process(name, fh):
            raise MemberFailure("unparsable DICOM")
        with self.assertRaisesRegex(FolderError, "none of the 1"):
            folder_bundle.run_per_file(tar, "dicom", self.path("out.tar"), process, lambda: None)
        self.assertFalse(os.path.exists(self.path("out.tar")))

    def test_image_output_names(self):
        self.assertEqual(folder_bundle.output_name("a/x.jpeg", "image", "/o/redacted.jpg"),
                         "a/x_redacted.jpg")
        self.assertEqual(folder_bundle.output_name("IM1", "dicom", "/o/after.dcm"), "IM1_anonymised.dcm")


class PerFileThroughTheSdk(TempDirCase):
    """P3DX_SDK._run_folder_per_file with the container replaced: staging of
    each member, workspace clearing, the status check, the handoff to
    finalize, and release of the bundle."""

    def setUp(self):
        super().setUp()
        import P3DX_SDK
        self.sdk = P3DX_SDK
        self.data, self.out, self.scratch = (os.path.join(self.dir, d) for d in ("data", "out", "scratch"))
        for d in (self.data, self.out, self.scratch):
            os.makedirs(d)
        for attr, value in (("tee_input_data", self.data), ("tee_output", self.out)):
            p = mock.patch.object(config.paths, attr, value)
            p.start()
            self.addCleanup(p.stop)
        self.staged_names = []

    def fake_container(self):
        staged = os.listdir(self.data)
        self.staged_names.extend(staged)
        assert len(staged) == 1, staged
        with open(os.path.join(self.data, staged[0]), "rb") as f:
            data = f.read()
        if data[128:132] != b"DICM":
            with open(os.path.join(self.out, "status.json"), "w") as f:
                json.dump({"status": "error", "error": "Unparsable DICOM: PatientName=DOE"}, f)
            return
        case = os.path.join(self.out, "input")
        os.makedirs(os.path.join(case, "keystore"))
        for name, body in (("after_deidentification.dcm", b"DEID" + data[-1:]),
                           ("before_deidentification.dcm", data),
                           ("keystore/fpe.json", b"{}")):
            with open(os.path.join(case, name), "wb") as f:
                f.write(body)

    def test_per_file_run(self):
        uid = "0f8fad5b-d9cb-469f-a165-70867728950e"
        tar = build_tar(os.path.join(self.scratch, f"{uid}.bin"),
                        [("s1/IM0001", DICOM + b"1"), ("s1/IM0002", b"garbage"), ("s2/IM0003", DICOM + b"3")])
        meta = {"filename": "scans.tar", "format": "dicom", "folder": {"mode": "per_file", "files_total": 3}}
        with open(os.path.join(self.scratch, f"{uid}.meta.json"), "w") as f:
            json.dump(meta, f)

        with mock.patch.object(self.sdk, "_run_application_once", self.fake_container), \
                mock.patch.object(self.sdk, "_direct_upload_id", return_value=uid), \
                mock.patch.object(direct_upload, "_scratch_dir", return_value=self.scratch):
            self.sdk.run_docker_containers()

        self.assertEqual(self.staged_names, ["input.dcm"] * 3, "each member alone, as .dcm")
        self.assertFalse(os.path.exists(tar), "bundle released after the loop")
        self.assertEqual(os.listdir(self.out), [self.sdk.FOLDER_OUTPUT_NAME])
        with tarfile.open(os.path.join(self.out, self.sdk.FOLDER_OUTPUT_NAME)) as tf:
            self.assertEqual(sorted(tf.getnames()), ["_manifest.json", "s1/IM0001_anonymised.dcm",
                                                     "s2/IM0003_anonymised.dcm"])
            manifest = json.load(tf.extractfile("_manifest.json"))
            blob = b"".join(tf.extractfile(m).read() for m in tf.getmembers())
        self.assertEqual(manifest["files"][1]["error"], "unparsable DICOM")
        self.assertNotIn(b"DOE", blob)
        self.assertNotIn(b"garbage", blob)
        stored = direct_upload.read_staged_meta(self.scratch, uid)
        self.assertEqual(stored["folder"], {"mode": "per_file", "files_total": 3,
                                            "files_succeeded": 2, "files_failed": 1})

    def test_a_single_file_run_is_untouched(self):
        called = []
        with mock.patch.object(self.sdk, "_run_application_once", lambda: called.append(1)), \
                mock.patch.object(self.sdk, "_direct_upload_id", return_value=None):
            self.sdk.run_docker_containers()
        self.assertEqual(called, [1])


DICOM_IMAGE = os.environ.get("SKALD_DICOM_IMAGE", "ghcr.io/datakaveri/skald-dicom:latest")
DICOM_SAMPLE = os.environ.get("SKALD_DICOM_SAMPLE",
                              os.path.expanduser("~/Downloads/DICOM/app/data/input.dcm"))


@unittest.skipUnless(os.environ.get("P3DX_DICOM_IT") == "1",
                     "set P3DX_DICOM_IT=1 to run the real SKALD-DICOM image (slow)")
class DicomConsistency(unittest.TestCase):
    """RECORDS current behaviour — it does not assert a requirement.

    Per-file mode runs the application once per member with a fresh output
    volume, so nothing (keystore, salt) carries from one member to the next.
    Observed on 2026-09-30 with skald-dicom:latest: the same PatientID in two
    members comes out with DIFFERENT hashes, under the default policy and under
    hashing_with_salt alike. If this test starts failing, cross-file
    consistency has arrived in the application; update
    backend-changes-direct-upload.md §2.9 and flip the assertion.

    Needs docker, the image, pydicom, and a sample DICOM (SKALD_DICOM_SAMPLE).
    The work directory is under $HOME because a snap-packaged docker cannot
    see /tmp.
    """

    def setUp(self):
        try:
            import pydicom
        except ImportError:
            self.skipTest("pydicom not installed")
        if not os.path.isfile(DICOM_SAMPLE):
            self.skipTest(f"no sample DICOM at {DICOM_SAMPLE}")
        self.pydicom = pydicom
        base = os.path.expanduser(os.environ.get("P3DX_DICOM_IT_DIR", "~/.cache/p3dx-dicom-it"))
        os.makedirs(base, exist_ok=True)
        self.dir = tempfile.mkdtemp(dir=base)
        self.addCleanup(self._cleanup)

    def _cleanup(self):
        # The container writes as root; remove its files the same way.
        subprocess.run(["docker", "run", "--rm", "--entrypoint", "rm", "-v", f"{self.dir}:/w",
                        DICOM_IMAGE, "-rf", "/w/data", "/w/output", "/w/config"], capture_output=True)
        shutil.rmtree(self.dir, ignore_errors=True)

    def patient_ids(self, job_config):
        members = []
        for i in range(2):
            ds = self.pydicom.dcmread(DICOM_SAMPLE)
            ds.PatientID, ds.PatientName = "PAT-12345", f"DOE^JOHN{i}"
            buf = io.BytesIO()
            ds.save_as(buf)
            members.append((f"s{i}/IM0001", buf.getvalue()))
        tar = build_tar(os.path.join(self.dir, "in.tar"), members)
        data, out, cfg = (os.path.join(self.dir, d) for d in ("data", "output", "config"))
        kept = os.path.join(self.dir, "kept")
        for d in (data, cfg, kept):
            os.makedirs(d, exist_ok=True)
        if job_config:
            with open(os.path.join(cfg, "config.json"), "w") as f:
                json.dump(job_config, f)

        def process(name, fh):
            subprocess.run(["docker", "run", "--rm", "--entrypoint", "rm", "-v", f"{self.dir}:/w",
                            DICOM_IMAGE, "-rf", "/w/output"], capture_output=True)
            os.makedirs(out, exist_ok=True)
            with open(os.path.join(data, "input.dcm"), "wb") as f:
                shutil.copyfileobj(fh, f)
            subprocess.run(["docker", "run", "--rm", "-v", f"{data}:/app/data", "-v", f"{cfg}:/app/config",
                            "-v", f"{out}:/app/output", DICOM_IMAGE], capture_output=True, timeout=1800)
            found = [os.path.join(r, n) for r, _, fs in os.walk(out) for n in fs
                     if n == "after_deidentification.dcm"]
            if len(found) != 1:
                raise MemberFailure("no de-identified output was produced")
            dst = os.path.join(kept, "out.dcm")
            shutil.copy(found[0], dst)
            return dst

        manifest = folder_bundle.run_per_file(tar, "dicom", os.path.join(self.dir, "out.tar"),
                                              process, lambda: None)
        self.assertEqual(manifest["files_succeeded"], 2)
        with tarfile.open(os.path.join(self.dir, "out.tar")) as tf:
            return [str(self.pydicom.dcmread(tf.extractfile(m)).PatientID)
                    for m in tf.getmembers() if m.name.endswith(".dcm")]

    def test_default_policy_same_patient_id_hashes_differ_across_files(self):
        a, b = self.patient_ids(None)
        self.assertNotIn("PAT-12345", (a, b))
        self.assertNotEqual(a, b, "cross-file consistency now holds — see the class docstring")

    def test_hashing_with_salt_same_patient_id_hashes_differ_across_files(self):
        a, b = self.patient_ids({"dicom_deidentify": {
            "operations": ["hashing_with_salt"],
            "tag_actions": {"PatientID": {"action": "hashing_with_salt"}}}})
        self.assertNotIn("PAT-12345", (a, b))
        self.assertNotEqual(a, b, "cross-file consistency now holds — see the class docstring")


# --------------------------------------------------------------------------- #
# 7. status.json
# --------------------------------------------------------------------------- #


class FolderStatus(unittest.TestCase):
    def test_blocks(self):
        self.assertEqual(folder_bundle.folder_status("joint", 250, records_in=250),
                         {"mode": "joint", "files_total": 250, "records_in": 250})
        self.assertEqual(folder_bundle.folder_status("per_file", 3, files_succeeded=2, files_failed=1),
                         {"mode": "per_file", "files_total": 3, "files_succeeded": 2, "files_failed": 1})

    def test_only_counts_survive_the_loopback_handoff(self):
        got = direct_upload._sanitise_folder_status(
            {"mode": "per_file", "files_total": 3, "files_succeeded": 2, "files_failed": 1,
             "files": [{"path": "secret"}]})
        self.assertEqual(got, {"mode": "per_file", "files_total": 3, "files_succeeded": 2, "files_failed": 1})
        self.assertIsNone(direct_upload._sanitise_folder_status({"mode": "zip"}))


# --------------------------------------------------------------------------- #
# Scratch: peak tmpfs use in both modes
# --------------------------------------------------------------------------- #


def _du(path):
    total = 0
    for root, _, files in os.walk(path):
        for n in files:
            try:
                total += os.path.getsize(os.path.join(root, n))
            except FileNotFoundError:
                pass
    return total


class ScratchPeak(TempDirCase):
    """Only the bundle and (per file) the output tar may live in scratch; each
    member's input and output go to data/ and output/. At the default size this
    runs a 30 MB bundle and bounds peak as a multiple of it; set
    P3DX_SCRATCH_FULL=1 to measure the real 300 MB case."""

    FULL = os.environ.get("P3DX_SCRATCH_FULL") == "1"
    BUNDLE = 290 * MB if FULL else 30 * MB
    LIMIT = 2 * 1024 * MB

    def test_per_file_peak(self):
        scratch = os.path.join(self.dir, "scratch")
        work = os.path.join(self.dir, "output")
        os.makedirs(scratch)
        os.makedirs(work)
        size = 3 * MB
        n = self.BUNDLE // size
        members = [(f"d{i // 50}/IM{i:05d}", DICOM + os.urandom(size - len(DICOM))) for i in range(n)]
        tar = build_tar(os.path.join(scratch, "u.bin"), members)
        del members
        bundle = os.path.getsize(tar)
        peak = [bundle]

        def process(name, fh):
            out = os.path.join(work, "after_deidentification.dcm")
            with open(out, "wb") as f:
                shutil.copyfileobj(fh, f)
            peak[0] = max(peak[0], _du(scratch))
            return out

        def cleanup():
            for x in os.listdir(work):
                os.unlink(os.path.join(work, x))
            peak[0] = max(peak[0], _du(scratch))

        folder_bundle.run_per_file(tar, "dicom", os.path.join(scratch, "u.out.tar"), process, cleanup)
        peak[0] = max(peak[0], _du(scratch))
        print(f"\n  per-file: bundle {bundle / MB:.0f} MB, peak scratch {peak[0] / MB:.0f} MB", flush=True)
        self.assertLessEqual(peak[0], 2.1 * bundle)
        self.assertLess(peak[0] * (300 * MB / bundle), self.LIMIT / 2,
                        "at 300 MB this would not stay well under 2 GiB")

    def test_joint_peak(self):
        scratch = os.path.join(self.dir, "scratch")
        data = os.path.join(self.dir, "data")
        os.makedirs(scratch)
        os.makedirs(data)
        row = "P{0:08d},{1},{2},flu,some free text to pad the row out a little\n"
        per_file = []
        chunk = 2 * MB
        written, i = 0, 0
        while written < self.BUNDLE:
            lines = ["patient_id,age,gender,diagnosis,notes\n"]
            size = len(lines[0])
            while size < chunk:
                line = row.format(i, 20 + i % 60, "MF"[i % 2])
                lines.append(line)
                size += len(line)
                i += 1
            per_file.append((f"f{len(per_file):04d}.csv", "".join(lines).encode()))
            written += size
        tar = build_tar(os.path.join(scratch, "u.bin"), per_file)
        del per_file
        bundle = os.path.getsize(tar)
        section = skald_section(quasi_identifiers={"numerical": [{"column": "age", "type": "int"}],
                                                   "categorical": ["gender"]},
                                hashing_with_salt=["patient_id"], sensitive_parameter="diagnosis",
                                insensitive_columns=["notes"])
        peak = [bundle]
        real_add = folder_bundle._Spool.add

        def add(spool, obj):
            if len(spool) % 50_000 == 0:
                peak[0] = max(peak[0], _du(scratch))
            real_add(spool, obj)

        with mock.patch.object(folder_bundle._Spool, "add", add):
            result = folder_bundle.combine_tabular(tar, "csv", section,
                                                   os.path.join(data, "p.csv"), data)
        peak[0] = max(peak[0], _du(scratch))
        print(f"\n  joint: bundle {bundle / MB:.0f} MB, {result.records_in:,} records, "
              f"peak scratch {peak[0] / MB:.0f} MB", flush=True)
        self.assertLessEqual(peak[0], 1.01 * bundle, "joint mode adds nothing to scratch")


if __name__ == "__main__":
    unittest.main()

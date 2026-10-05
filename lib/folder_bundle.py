"""
Folder inputs. A folder reaches the enclave one of two ways:

- uploaded, as ONE uncompressed POSIX tar through the unchanged direct-upload
  transport (job -> /upload/init -> chunks -> complete). Decryption, AAD,
  digests and the digest root all apply to the tar bytes exactly as to any
  other file; nothing here touches them.
- in blob mode, as a blob PREFIX (lib/blob_folder.py), each blob encrypted like
  a single-blob input.

Either way the pipelines below see the same thing: a sequence of `Member`s.

Two modes, chosen by the member format (every member is of `format`):

- csv / json / excel -> JOINT. All members are one dataset. They are combined
  into ONE input file of the member format, in data/, and SKALD runs once over
  it, so k-anonymity holds across the whole folder and one run means one salt
  and one key for every file.
- dicom / image -> PER FILE. Each member runs through the existing single-file
  path on its own; the outputs come back as `<folder>_anonymised.tar` with a
  root `_manifest.json`.

Why joint mode builds a combined FILE rather than a combined table: SKALD's
loaders live inside its container (a Rust binary), not in this process. Handing
it one file in the member format is how its own loader stays the one that
parses the data. For JSON that is exact: SKALD's reader takes an array of
records and unions their keys, so the combined array of every member's records
is read the same way each member would be. CSV members are re-emitted over the
union of columns. Excel is the one format that has to be merged here, because
each workbook's sheets must be joined BEFORE the workbooks are stacked; that
merge is a port of SKALD's own (multitabular.rs merge_excel_sheets), and the
result is handed to SKALD as a single-sheet workbook, which it reads as-is.

This module has no Flask or docker dependency. lib/direct_upload.py calls the
validation at bundle-upload time; Fetch_data/fetch_data.py calls the combiner;
P3DX_SDK.py drives the per-file loop, whose output streams into a sink
(lib/output_stream.py) rather than being staged.
"""

from __future__ import annotations

import csv
import datetime as _dt
import io
import json
import math
import os
import posixpath
import secrets
import tarfile
from array import array
from collections import defaultdict
from dataclasses import dataclass, field
from typing import BinaryIO, Callable, Iterator

from enclave.enclave_direct_upload import (
    MAX_FOLDER_BYTES_BY_FORMAT,
    MAX_FOLDER_FILES,
    MAX_TOTAL_BYTES_BY_FORMAT,
    MB,
    is_folder_filename as is_folder_name,
)

MEMBER_FORMATS = ("csv", "json", "excel", "dicom", "image")
JOINT_FORMATS = ("csv", "json", "excel")
PER_FILE_FORMATS = ("dicom", "image")

MODE_JOINT = "joint"
MODE_PER_FILE = "per_file"

#: Offset and value of the POSIX ustar magic in a tar header.
USTAR_OFFSET = 257
USTAR_MAGIC = b"ustar"

MANIFEST_NAME = "_manifest.json"

#: Data rows one worksheet can hold (1,048,576 rows, less the header). An Excel
#: folder combines into ONE sheet, and SKALD writes its result as one sheet, so
#: past this neither the input nor the output can exist.
EXCEL_MAX_ROWS = 1_048_576 - 1
DP_REJECTION = "folder upload is not supported for differential privacy yet"

_BLOCK = 512
_APPLICATION_BY_FORMAT = {"dicom": "skald_dicom", "image": "skald_image"}


class FolderError(Exception):
    """Fails the whole job. The message leaves the enclave (status.json and the
    bundle-upload response), so it may name a member path — the user's own —
    but never any member's contents."""


class MemberFailure(Exception):
    """Fails one member of a per-file folder. `str()` goes into the output
    manifest verbatim, so raise it only with a fixed, content-free message."""


def mode_for(fmt: str) -> str:
    return MODE_JOINT if fmt in JOINT_FORMATS else MODE_PER_FILE


@dataclass
class Member:
    """One file of a folder. `open()` returns a binary file (usable as a
    context manager), or raises MemberFailure with a content-free reason when
    the member cannot be produced — e.g. a blob that changed after listing."""
    name: str
    open: Callable[[], BinaryIO]


def tar_members(tar_path: str) -> Iterator[Member]:
    """The members of an uploaded (already validated) folder tar, read in place."""
    with tarfile.open(tar_path, "r:") as tf:
        for m in tf:
            yield Member(m.name, lambda m=m: tf.extractfile(m))


# --------------------------------------------------------------------------- #
# 1. Detection
# --------------------------------------------------------------------------- #


def has_ustar_magic(path: str) -> bool:
    try:
        with open(path, "rb") as fh:
            fh.seek(USTAR_OFFSET)
            return fh.read(len(USTAR_MAGIC)) == USTAR_MAGIC
    except OSError:
        return False


def detect(filename, path: str, fmt: str) -> bool:
    """True for a folder job: a `.tar` name AND the ustar magic. Anything else
    is a single-file job, handled exactly as before.

    Raises FolderError — before any parser runs — for a `.tar` name without the
    magic, or the magic under a format that is not a member format."""
    named_tar = is_folder_name(filename)
    magic = has_ustar_magic(path) if named_tar else False
    if named_tar and not magic:
        raise FolderError("the upload is named .tar but is not a tar archive")
    if not named_tar:
        return False
    if (fmt or "").lower() not in MEMBER_FORMATS:
        raise FolderError(
            f"format '{fmt}' is not a folder member format; expected one of "
            f"{', '.join(MEMBER_FORMATS)}"
        )
    return True


# --------------------------------------------------------------------------- #
# 3. Validation — one walk over the archive, before any pipeline starts
# --------------------------------------------------------------------------- #


@dataclass
class ArchiveSummary:
    fmt: str
    mode: str
    members: list = field(default_factory=list)   # [(path, size)], archive order
    total_bytes: int = 0

    @property
    def files_total(self) -> int:
        return len(self.members)


def _raw_header_flags(fh: BinaryIO, start: int, end: int) -> list:
    """Typeflags of every raw header from `start` (the first extension header
    tarfile folded into this member) up to the member's own header.

    tarfile silently folds GNU long-name ('L'/'K') and pax ('x'/'g') headers
    into the member it returns, so a GNU long name is invisible on the TarInfo.
    Walking the raw headers is the only way to reject it."""
    flags = []
    pos = start
    while True:
        fh.seek(pos)
        buf = fh.read(_BLOCK)
        if len(buf) < _BLOCK:
            raise FolderError("the folder archive is truncated")
        flags.append(buf[156:157])
        if pos + _BLOCK >= end:
            return flags
        try:
            size = tarfile.nti(buf[124:136])
        except tarfile.InvalidHeaderError:
            raise FolderError("the folder archive has a malformed header") from None
        pos += _BLOCK + -(-size // _BLOCK) * _BLOCK
        if pos >= end:
            raise FolderError("the folder archive has a malformed header")


def _check_path(name: str) -> None:
    if not name or "\x00" in name:
        raise FolderError("the folder archive has a member with an empty or NUL-bearing path")
    if name.startswith("/") or "\\" in name or (len(name) > 1 and name[1] == ":"):
        raise FolderError(f"folder member '{_printable(name)}' has an absolute path")
    parts = name.split("/")
    if ".." in parts:
        raise FolderError(f"folder member '{_printable(name)}' has a '..' path component")
    if posixpath.normpath(name) != name or "." in parts:
        raise FolderError(f"folder member '{_printable(name)}' has a non-canonical path")
    root = "/folder-root"
    if not posixpath.normpath(posixpath.join(root, name)).startswith(root + "/"):
        raise FolderError(f"folder member '{_printable(name)}' resolves outside the folder")


def _printable(name: str) -> str:
    return name.encode("unicode_escape").decode("ascii")[:200]


def validate_archive(path: str, fmt: str,
                     head_matches: Callable[[bytes, str], bool]) -> ArchiveSummary:
    """Walk every member once and fail the whole job on anything the UI cannot
    have produced: a non-regular member, an unsafe or duplicate path, a count or
    size over the limits, an empty archive, or a member whose content does not
    sniff as `fmt`. Nothing is extracted; each member's head is read in place.

    `head_matches` is the existing single-file sniff (lib.direct_upload
    .head_matches_format), passed in so both paths share one definition."""
    fmt = fmt.lower()
    per_member_cap = MAX_TOTAL_BYTES_BY_FORMAT[fmt]
    folder_cap = MAX_FOLDER_BYTES_BY_FORMAT[fmt]
    summary = ArchiveSummary(fmt=fmt, mode=mode_for(fmt))
    seen = set()
    file_size = os.path.getsize(path)
    end_of_members = 0

    try:
        with open(path, "rb") as raw, tarfile.open(path, "r:") as tf:
            for member in tf:
                if len(summary.members) >= MAX_FOLDER_FILES:
                    raise FolderError(
                        f"the folder has more than {MAX_FOLDER_FILES:,} files"
                    )
                name = member.name
                if member.type != tarfile.REGTYPE or member.sparse is not None:
                    raise FolderError(
                        f"folder member '{_printable(name)}' is not a regular file"
                    )
                flags = _raw_header_flags(raw, member.offset, member.offset_data)
                if flags[-1] != tarfile.REGTYPE or any(f != tarfile.XHDTYPE for f in flags[:-1]):
                    raise FolderError(
                        f"folder member '{_printable(name)}' uses an unsupported "
                        f"tar header type"
                    )
                _check_path(name)
                if name in seen:
                    raise FolderError(f"folder member '{_printable(name)}' appears twice")
                seen.add(name)

                if member.size <= 0:
                    raise FolderError(f"folder member '{_printable(name)}' is empty")
                if member.size > per_member_cap:
                    raise FolderError(
                        f"folder member '{_printable(name)}' is over the "
                        f"{per_member_cap // MB} MB limit for {fmt.upper()} files"
                    )
                summary.total_bytes += member.size
                if summary.total_bytes > folder_cap:
                    raise FolderError(
                        f"the folder's files add up to more than the {folder_cap // MB} MB "
                        f"limit for a {fmt.upper()} folder"
                    )
                data_end = member.offset_data + member.size
                if data_end > file_size:
                    raise FolderError("the folder archive is truncated")
                end_of_members = member.offset_data + -(-member.size // _BLOCK) * _BLOCK

                with tf.extractfile(member) as fh:
                    head = fh.read(512)
                if not head_matches(head, fmt):
                    raise FolderError(
                        f"folder member '{_printable(name)}' does not match the "
                        f"declared format '{fmt}'"
                    )
                summary.members.append((name, member.size))

            # tarfile ends iteration quietly at the first bad header after the
            # first member, which would leave whatever follows unexamined. The
            # archive must end in zero blocks and nothing else.
            raw.seek(end_of_members)
            while True:
                block = raw.read(1 << 16)
                if not block:
                    break
                if block.strip(b"\x00"):
                    raise FolderError("the folder archive has data after its last file")
    except tarfile.TarError:
        raise FolderError("the folder archive is not a readable tar file") from None

    if not summary.members:
        raise FolderError("the folder archive is empty")
    return summary


# --------------------------------------------------------------------------- #
# Config — the settings live under the key `data_type` names (`<folder>`)
# --------------------------------------------------------------------------- #


def load_section(config_dir: str) -> tuple:
    """(data_type, section) from the app config the bundle delivered."""
    try:
        names = sorted(os.listdir(config_dir))
    except OSError as exc:
        raise FolderError(f"could not read the job configuration: {exc}") from None
    for name in names:
        if not name.endswith(".json"):
            continue
        try:
            with open(os.path.join(config_dir, name)) as f:
                cfg = json.load(f)
        except (OSError, ValueError):
            continue
        if not isinstance(cfg, dict):
            continue
        data_type = cfg.get("data_type")
        if isinstance(data_type, str) and isinstance(cfg.get(data_type), dict):
            return data_type, cfg[data_type]
    raise FolderError("the job configuration has no section for its data_type")


#: Keys whose value is a list of column names.
_COLUMN_LIST_KEYS = (
    "insensitive_columns", "suppress", "hashing_with_salt", "hashing_without_salt",
    "charcloak", "sensitive_columns",
)
#: Keys whose value is an object keyed by column name.
_COLUMN_KEYED_KEYS = ("size", "qi_constraints", "fixed_bins", "categorical_hierarchies")


def configured_columns(section: dict) -> set:
    """Every column the config names anywhere, by technique or as insensitive.

    The UI writes every column of the (union) dataset into
    `insensitive_columns` up front and moves them out as techniques are chosen,
    so a column that is in the data but named nowhere here was not in the
    folder the user configured."""
    cols = set()

    def walk(node):
        if isinstance(node, dict):
            col = node.get("column")
            if isinstance(col, str):
                cols.add(col)
            for key, value in node.items():
                if isinstance(value, list) and (key in _COLUMN_LIST_KEYS or key.endswith("columns")):
                    cols.update(v for v in value if isinstance(v, str))
                walk(value)
        elif isinstance(node, list):
            for item in node:
                walk(item)

    walk(section)
    for key in _COLUMN_KEYED_KEYS:
        value = section.get(key)
        if isinstance(value, dict):
            cols.update(value.keys())
    sensitive = section.get("sensitive_parameter")
    if isinstance(sensitive, str):
        cols.add(sensitive)
    elif isinstance(sensitive, list):
        cols.update(v for v in sensitive if isinstance(v, str))
    qis = section.get("quasi_identifiers")
    if isinstance(qis, dict):
        for c in qis.get("categorical") or []:
            if isinstance(c, str):
                cols.add(c)
    return cols


def quasi_identifiers(section: dict) -> tuple:
    """({numerical column: 'int'|'float'}, {categorical column})."""
    numerical, categorical = {}, set()
    qis = section.get("quasi_identifiers")
    if isinstance(qis, dict):
        for n in qis.get("numerical") or []:
            if isinstance(n, dict) and isinstance(n.get("column"), str):
                numerical[n["column"]] = str(n.get("type") or "int").lower()
        for c in qis.get("categorical") or []:
            if isinstance(c, dict) and isinstance(c.get("column"), str):
                categorical.add(c["column"])
            elif isinstance(c, str):
                categorical.add(c)
    return numerical, categorical


# --------------------------------------------------------------------------- #
# 4. Joint: load every member, combine, write ONE input for SKALD
# --------------------------------------------------------------------------- #


class _Spool:
    """Records on disk, one JSON line each, so memory holds one member at a
    time plus an offset per record — not the whole folder. `shuffled()` reads
    them back in a cryptographically random order."""

    def __init__(self, path: str):
        self.path = path
        self._fh = open(path, "w+b")
        os.chmod(path, 0o600)
        self._offsets = array("Q")

    def __len__(self):
        return len(self._offsets)

    def add(self, obj) -> None:
        self._offsets.append(self._fh.tell())
        self._fh.write(json.dumps(obj, ensure_ascii=False, separators=(",", ":"),
                                  default=_encode_cell).encode("utf-8"))
        self._fh.write(b"\n")

    def shuffled(self) -> Iterator:
        self._fh.flush()
        order = list(range(len(self._offsets)))
        secrets.SystemRandom().shuffle(order)
        for i in order:
            self._fh.seek(self._offsets[i])
            yield json.loads(self._fh.readline(), object_hook=_decode_cell)

    def close(self) -> None:
        try:
            self._fh.close()
        finally:
            try:
                os.unlink(self.path)
            except FileNotFoundError:
                pass


def _encode_cell(value):
    # Excel cells keep their type through the spool so the combined workbook
    # carries a date as a date, and SKALD's own reader formats it.
    for kind, cls in (("datetime", _dt.datetime), ("date", _dt.date), ("time", _dt.time)):
        if isinstance(value, cls):
            return {"$cell": kind, "v": value.isoformat()}
    if isinstance(value, _dt.timedelta):
        return {"$cell": "timedelta", "v": value.total_seconds()}
    raise TypeError(f"cannot spool a {type(value).__name__}")


def _decode_cell(obj):
    kind = obj.get("$cell") if len(obj) == 2 else None
    if kind == "datetime":
        return _dt.datetime.fromisoformat(obj["v"])
    if kind == "date":
        return _dt.date.fromisoformat(obj["v"])
    if kind == "time":
        return _dt.time.fromisoformat(obj["v"])
    if kind == "timedelta":
        return _dt.timedelta(seconds=obj["v"])
    return obj


def _json_scalar_to_string(value) -> str:
    """SKALD's json_scalar_to_string — what a JSON value becomes as a cell."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, (int, float, str)):
        return str(value) if not isinstance(value, float) else repr(value)
    return json.dumps(value, separators=(",", ":"))


def _is_number(text: str) -> bool:
    s = text.strip()
    if "_" in s:          # Python accepts 1_000; SKALD's parse::<f64> does not
        return False
    try:
        return math.isfinite(float(s))
    except ValueError:
        return False


def _load_json_member(fh: BinaryIO, name: str) -> Iterator[dict]:
    """A member's records. Top-level OBJECT: one record, even when some fields
    are arrays (a patient with a `visits` list is one patient). Top-level
    array: a list of records, as SKALD reads a single file. The UI previews a
    folder with this same rule (headPreview objectIsRecord)."""
    try:
        text = fh.read().decode("utf-8-sig")
    except UnicodeDecodeError:
        raise FolderError(f"folder member '{_printable(name)}' is not UTF-8 JSON") from None

    def reject_constant(token):
        raise ValueError(f"non-standard JSON value {token}")

    try:
        value = json.loads(text, parse_constant=reject_constant)
    except ValueError as exc:
        raise FolderError(
            f"folder member '{_printable(name)}' is not valid JSON "
            f"(line {getattr(exc, 'lineno', '?')}, column {getattr(exc, 'colno', '?')})"
        ) from None
    del text
    if isinstance(value, dict):
        records = [value]
    elif isinstance(value, list):
        if not value:
            raise FolderError(f"folder member '{_printable(name)}' holds no records")
        records = value
    else:
        raise FolderError(f"folder member '{_printable(name)}' holds a single value, not a record")
    for i, rec in enumerate(records):
        if not isinstance(rec, dict):
            raise FolderError(
                f"folder member '{_printable(name)}': record {i + 1} is not a JSON object"
            )
        for v in rec.values():
            if isinstance(v, float) and not math.isfinite(v):
                raise FolderError(f"folder member '{_printable(name)}' has a number out of range")
        yield rec


def _load_csv_member(fh: BinaryIO, name: str) -> Iterator[tuple]:
    """(header, row) pairs. Rows shorter than the header are padded with empty
    cells; a row LONGER than the header has values with no column to put them
    in and fails the member, since they could otherwise pass unconfigured."""
    text = io.TextIOWrapper(fh, encoding="utf-8-sig", newline="")
    reader = csv.reader(text)
    try:
        header = next(reader, None)
        if not header:
            raise FolderError(f"folder member '{_printable(name)}' has no header row")
        if len(set(header)) != len(header):
            raise FolderError(f"folder member '{_printable(name)}' repeats a column name")
        width = len(header)
        for row in reader:
            if not row or (len(row) == 1 and not row[0].strip()):
                continue
            if len(row) > width:
                raise FolderError(
                    f"folder member '{_printable(name)}': row {reader.line_num} has "
                    f"more values than the header has columns"
                )
            if len(row) < width:
                row = row + [""] * (width - len(row))
            yield header, row
    except UnicodeDecodeError:
        raise FolderError(f"folder member '{_printable(name)}' is not UTF-8 text") from None
    except csv.Error as exc:
        raise FolderError(
            f"folder member '{_printable(name)}' is not valid CSV (line {reader.line_num})"
        ) from None
    finally:
        text.detach()


@dataclass
class JointResult:
    path: str
    files_total: int
    records_in: int
    columns: list


def combine_tabular(members, fmt: str, section: dict, out_path: str,
                    spool_dir: str) -> JointResult:
    """Load every member with its format's loader, concatenate over the union
    of columns, and write ONE shuffled input file of the member format.

    `members` is an iterable of Member (tar_members, or a blob listing). Each
    member's raw bytes are dropped as soon as it is loaded.

    Fails the whole job — naming the member — on any member that cannot be
    loaded or produced: dropping one quietly changes the dataset k is computed
    over, and with one patient per file a patient would vanish from the release.

    Rows are shuffled HERE, with secrets.SystemRandom, as they are written for
    SKALD. Input order follows the sorted file paths, and SKALD keeps row
    order, so without this the output rows would map straight back to file
    names. Shuffling on the way in gives the output that same random order and
    works for every format, including a workbook, without re-reading SKALD's
    result. No source-file column is added at any stage.
    """
    fmt = fmt.lower()
    numerical, categorical = quasi_identifiers(section)
    qi_columns = set(numerical) | categorical
    configured = configured_columns(section)
    sheet_joins = _parse_sheet_joins(section) if fmt == "excel" else []

    columns, seen = [], set()
    missing_qi = defaultdict(list)       # column -> member paths lacking it
    files = 0
    spool = _Spool(os.path.join(spool_dir, ".folder-records.spool"))

    def note_columns(cols):
        for c in cols:
            if c not in seen:
                seen.add(c)
                columns.append(c)

    def check_record(name, keys, value_of):
        for qi in qi_columns:
            if qi not in keys:
                if not missing_qi[qi] or missing_qi[qi][-1] != name:
                    missing_qi[qi].append(name)
        for col, dtype in numerical.items():
            if col in keys:
                text = value_of(col)
                if text.strip() and not _is_number(text):
                    raise FolderError(
                        f"folder member '{_printable(name)}': a value in quasi-identifier "
                        f"column '{col}' is not a valid {dtype}"
                    )

    try:
        for member in members:
            files += 1
            name = member.name
            try:
                fh = member.open()
            except MemberFailure as exc:
                raise FolderError(f"folder member '{_printable(name)}': {exc}") from None
            with fh:
                if fmt == "json":
                    for rec in _load_json_member(fh, name):
                        note_columns(rec.keys())
                        check_record(name, rec, lambda c, r=rec: _json_scalar_to_string(r[c]))
                        spool.add(rec)
                elif fmt == "csv":
                    for header, row in _load_csv_member(fh, name):
                        rec = dict(zip(header, row))
                        note_columns(header)
                        check_record(name, rec, rec.__getitem__)
                        spool.add(rec)
                else:
                    sheet = _load_excel_member(fh, name, sheet_joins)
                    if len(set(sheet.columns)) != len(sheet.columns):
                        raise FolderError(
                            f"folder member '{_printable(name)}' has a repeated column "
                            f"name after its sheets are joined"
                        )
                    note_columns(sheet.columns)
                    if len(spool) + len(sheet.rows) > EXCEL_MAX_ROWS:
                        raise FolderError(
                            f"the folder's workbooks add up to more than {EXCEL_MAX_ROWS:,} "
                            f"rows (reached at '{_printable(name)}'), the most one Excel "
                            f"sheet can hold; the combined dataset and its output are each "
                            f"one sheet"
                        )
                    for row in sheet.rows:
                        rec = dict(zip(sheet.columns, row))
                        check_record(name, rec, lambda c, r=rec: _cell_str(r[c]))
                        spool.add(rec)

        unknown = [c for c in columns if c not in configured]
        if unknown:
            raise FolderError(
                "the folder has column(s) the configuration never mentions: "
                f"{unknown}. The files may have changed after they were selected, "
                "or a workbook differs from the sample. Every column needs a "
                "technique or an explicit 'insensitive' choice."
            )
        if missing_qi:
            detail = "; ".join(
                f"'{col}' is missing from {len(paths)} file(s), first "
                f"'{_printable(paths[0])}'" for col, paths in sorted(missing_qi.items())
            )
            raise FolderError(
                f"some records lack a quasi-identifier column ({detail}). Combining "
                "the folder would leave those values empty, and SKALD has no "
                "defined handling for an empty quasi-identifier, so the job is "
                "refused rather than guessing."
            )
        if not len(spool):
            raise FolderError("the folder holds no records")

        _write_combined(fmt, columns, spool, out_path)
        return JointResult(out_path, files, len(spool), columns)
    finally:
        spool.close()


def _write_combined(fmt: str, columns: list, spool: _Spool, out_path: str) -> None:
    tmp = out_path + ".partial"
    try:
        if fmt == "json":
            # Each member's records verbatim: SKALD's JSON reader unions the
            # keys itself and fills a missing one with "", exactly as it does
            # within a single file.
            with open(tmp, "w", encoding="utf-8") as f:
                f.write("[\n")
                for i, rec in enumerate(spool.shuffled()):
                    if i:
                        f.write(",\n")
                    f.write(json.dumps(rec, ensure_ascii=False))
                f.write("\n]\n")
        elif fmt == "csv":
            with open(tmp, "w", encoding="utf-8", newline="") as f:
                w = csv.writer(f, lineterminator="\n")
                w.writerow(columns)
                for rec in spool.shuffled():
                    w.writerow([rec.get(c, "") for c in columns])
        else:
            openpyxl = _openpyxl()
            wb = openpyxl.Workbook(write_only=True)
            ws = wb.create_sheet("Sheet1")
            ws.append(columns)
            for rec in spool.shuffled():
                ws.append([rec.get(c) for c in columns])
            wb.save(tmp)
        os.chmod(tmp, 0o600)
        os.replace(tmp, out_path)
    except BaseException:
        try:
            os.unlink(tmp)
        except FileNotFoundError:
            pass
        raise


# ── Excel: a port of SKALD's merge_excel_sheets (multitabular.rs) ─────────── #
#
# Each workbook is merged exactly as SKALD merges a single workbook, THEN the
# workbooks are stacked. Stacking first and letting SKALD join would match rows
# ACROSS workbooks wherever two of them reuse a join-key value.


@dataclass
class _Sheet:
    columns: list
    rows: list


@dataclass
class _Join:
    left: str
    right: str
    on: list
    how: str


_ALLOWED_HOW = ("left", "right", "inner", "outer", "cross")


def _openpyxl():
    try:
        import openpyxl  # noqa: PLC0415 — only Excel folders need it
    except ImportError:
        raise FolderError(
            "Excel folder uploads need the openpyxl package on the enclave host"
        ) from None
    return openpyxl


def _parse_sheet_joins(section: dict) -> list:
    """SKALD's parse_sheet_joins. The UI's `join_keys` is not read: SKALD does
    not read it for a single workbook either (it auto-joins on shared columns
    when `sheet_joins` is absent), and a folder must merge the way one file
    does."""
    joins = []
    for i, entry in enumerate(section.get("sheet_joins") or [], start=1):
        if not isinstance(entry, dict) or not isinstance(entry.get("left"), str) \
                or not isinstance(entry.get("right"), str):
            raise FolderError(f"sheet_joins step {i} needs 'left' and 'right' sheet names")
        on = entry.get("on")
        on = [on] if isinstance(on, str) else [c for c in (on or []) if isinstance(c, str)]
        if not on:
            raise FolderError(f"sheet_joins step {i} must list at least one join column")
        how = entry.get("how") or "left"
        if how not in _ALLOWED_HOW:
            raise FolderError(f"sheet_joins step {i}: 'how' must be one of {list(_ALLOWED_HOW)}")
        joins.append(_Join(entry["left"], entry["right"], on, how))
    return joins


def _cell_str(value) -> str:
    """calamine's cell_to_string, used for join keys and header names. The cell
    VALUES keep their type into the combined workbook, so SKALD's own reader is
    what finally formats them."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return value.isoformat(sep=" ") if isinstance(value, _dt.datetime) else value.isoformat()
    return str(value)


def _load_excel_member(fh: BinaryIO, name: str, joins: list) -> _Sheet:
    head = fh.read(8)
    fh.seek(0)
    if head.startswith(b"\xD0\xCF\x11\xE0"):
        raise FolderError(
            f"folder member '{_printable(name)}' is a legacy .xls workbook; folder "
            f"uploads read .xlsx only"
        )
    openpyxl = _openpyxl()
    try:
        wb = openpyxl.load_workbook(fh, read_only=True, data_only=True)
    except Exception:
        raise FolderError(f"folder member '{_printable(name)}' is not a readable workbook") from None
    try:
        sheets = [(ws.title, _range_to_sheet(ws.iter_rows(values_only=True))) for ws in wb.worksheets]
    except Exception:
        raise FolderError(f"folder member '{_printable(name)}' is not a readable workbook") from None
    finally:
        wb.close()
    if not sheets:
        raise FolderError(f"folder member '{_printable(name)}' has no sheets")
    return _merge_excel_sheets(sheets, joins, name)


def _range_to_sheet(rows) -> _Sheet:
    """calamine's range covers the bounding box of non-empty cells; the first
    row of it is the header, and every row is padded/truncated to its width."""
    grid, r0, r1, c0, c1 = [], None, None, None, None
    for i, row in enumerate(rows):
        row = list(row)
        grid.append(row)
        cols = [j for j, v in enumerate(row) if v is not None]
        if cols:
            r0 = i if r0 is None else r0
            r1 = i
            c0 = cols[0] if c0 is None else min(c0, cols[0])
            c1 = cols[-1] if c1 is None else max(c1, cols[-1])
    if r0 is None:
        return _Sheet([], [])
    box = [(r + [None] * (c1 + 1))[c0:c1 + 1] for r in grid[r0:r1 + 1]]
    return _Sheet([_cell_str(v) for v in box[0]], box[1:])


def _merge_excel_sheets(sheets: list, joins: list, source: str) -> _Sheet:
    if len(sheets) == 1:
        return sheets[0][1]
    if joins:
        return _apply_sheet_joins(sheets, joins, source)
    first = sorted(sheets[0][1].columns)
    if all(sorted(s.columns) == first for _, s in sheets):
        return _vertical_concat(sheets)
    return _auto_join_sheets(sheets)


def _key(row, idx):
    return tuple(_cell_str(row[i]) for i in idx)


def _merge_two(left: _Sheet, right: _Sheet, on: list, how: str) -> _Sheet:
    if how == "cross":
        return _cross_join(left, right)
    l_on = [left.columns.index(c) for c in on]
    r_on = [right.columns.index(c) for c in on]
    r_nonkey = [i for i in range(len(right.columns)) if i not in r_on]
    r_names = {right.columns[i] for i in r_nonkey}
    l_names = set(left.columns)
    out_cols = [f"{c}_x" if c in r_names else c for c in left.columns]
    out_cols += [f"{right.columns[i]}_y" if right.columns[i] in l_names else right.columns[i]
                 for i in r_nonkey]

    index = defaultdict(list)
    for ridx, row in enumerate(right.rows):
        index[_key(row, r_on)].append(ridx)

    out_rows, matched = [], set()
    for lrow in left.rows:
        matches = index.get(_key(lrow, l_on))
        if matches:
            for ridx in matches:
                matched.add(ridx)
                out_rows.append(list(lrow) + [right.rows[ridx][i] for i in r_nonkey])
        elif how in ("left", "outer"):
            out_rows.append(list(lrow) + [None] * len(r_nonkey))
    if how in ("right", "outer"):
        for ridx, rrow in enumerate(right.rows):
            if ridx in matched:
                continue
            row = [None] * len(left.columns)
            for pos, i in enumerate(l_on):
                row[i] = rrow[r_on[pos]]
            out_rows.append(row + [rrow[i] for i in r_nonkey])
    return _Sheet(out_cols, out_rows)


def _cross_join(left: _Sheet, right: _Sheet) -> _Sheet:
    r_names, l_names = set(right.columns), set(left.columns)
    cols = [f"{c}_x" if c in r_names else c for c in left.columns]
    cols += [f"{c}_y" if c in l_names else c for c in right.columns]
    return _Sheet(cols, [list(l) + list(r) for l in left.rows for r in right.rows])


def _apply_sheet_joins(sheets: list, joins: list, source: str) -> _Sheet:
    frames = dict(sheets)
    root = None
    for step_no, step in enumerate(joins, start=1):
        for side in (step.left, step.right):
            if side not in frames:
                raise FolderError(
                    f"folder member '{_printable(source)}': sheet_joins step {step_no} "
                    f"names sheet '{side}', which this workbook does not have"
                )
        left, right = frames[step.left], frames[step.right]
        missing = [c for c in step.on if c not in left.columns or c not in right.columns]
        if missing:
            raise FolderError(
                f"folder member '{_printable(source)}': sheet_joins step {step_no} join "
                f"column(s) {missing} are missing from a sheet"
            )
        frames[step.left] = _merge_two(left, right, step.on, step.how)
        root = root or step.left
    return frames[root]


def _vertical_concat(sheets: list) -> _Sheet:
    columns = sheets[0][1].columns
    rows = []
    for _, sheet in sheets:
        if sheet.columns == columns:
            rows.extend(sheet.rows)
        else:
            idx = [sheet.columns.index(c) for c in columns]
            rows.extend([r[i] for i in idx] for r in sheet.rows)
    return _Sheet(list(columns), rows)


def _auto_join_sheets(sheets: list) -> _Sheet:
    acc = sheets[0][1]
    for _, right in sheets[1:]:
        common = [c for c in acc.columns if c in right.columns]
        acc = _merge_two(acc, right, common, "left") if common else _horizontal_concat(acc, right)
    return acc


def _horizontal_concat(left: _Sheet, right: _Sheet) -> _Sheet:
    n = max(len(left.rows), len(right.rows))
    rows = []
    for i in range(n):
        l = list(left.rows[i]) if i < len(left.rows) else [None] * len(left.columns)
        r = list(right.rows[i]) if i < len(right.rows) else [None] * len(right.columns)
        rows.append(l + r)
    return _Sheet(left.columns + right.columns, rows)


# --------------------------------------------------------------------------- #
# 5. Per file: DICOM and image, streaming one member at a time
# --------------------------------------------------------------------------- #

_IMAGE_EXT_BY_MAGIC = (
    (b"\x89PNG\r\n\x1a\n", ".png"),
    (b"\xff\xd8\xff", ".jpg"),
    (b"BM", ".bmp"),
    (b"II*\x00", ".tif"),
    (b"MM\x00*", ".tif"),
)


def input_extension(fmt: str, head: bytes) -> str:
    """The extension a member is staged under for its container. Chosen from
    the content, so a DICOM member with no extension (`IM0001`) still reads as
    .dcm and an image keeps an extension the image app accepts."""
    if fmt == "dicom":
        return ".dcm"
    if head[:4] == b"RIFF" and head[8:12] == b"WEBP":
        return ".webp"
    for magic, ext in _IMAGE_EXT_BY_MAGIC:
        if head.startswith(magic):
            return ext
    return ".png"


def output_name(member_name: str, fmt: str, produced_path: str) -> str:
    """The member's usual single-file output name, at its relative path:
    `s1/IM0001` -> `s1/IM0001_anonymised.dcm`, `a/x.png` -> `a/x_redacted.png`
    (see P3DX_SDK._upload_direct_output for the single-file names)."""
    directory, base = posixpath.split(member_name)
    stem = os.path.splitext(base)[0] or base
    if fmt == "dicom":
        out = f"{stem}_anonymised.dcm"
    else:
        out = f"{stem}_redacted{os.path.splitext(produced_path)[1].lower()}"
    return posixpath.join(directory, out) if directory else out


def run_per_file(members, fmt: str, sink,
                 process_member: Callable[[str, BinaryIO], str],
                 cleanup_member: Callable[[], None]) -> tuple:
    """De-identify each member on its own and stream each result into `sink`
    the moment it exists, one member on disk at a time.

    `members` is an iterable of Member. `process_member(name, fileobj)` runs the
    existing single-file path and returns the path of the output it produced;
    `cleanup_member()` removes that member's input and output whatever
    happened. `sink` takes `add_file(arcname, path)`, `close(manifest_name,
    manifest_bytes)` and `abort()` (lib/output_stream.StreamingTarSink, or the
    loopback sink that drives one in the enclave manager).

    A member that fails — including one whose source could not be produced —
    is LEFT OUT of the output, never copied through raw, and listed in the
    manifest with a fixed, content-free reason. Returns (manifest, whatever
    sink.close returned); raises FolderError, after aborting the sink, if no
    member succeeded.
    """
    results, used = [], set()
    try:
        for member in members:
            entry = {"path": member.name}
            try:
                try:
                    with member.open() as fh:
                        produced = process_member(member.name, fh)
                    arcname = output_name(member.name, fmt, produced)
                    if arcname in used or arcname == MANIFEST_NAME:
                        raise MemberFailure("output name collides with another file's output")
                except MemberFailure as exc:
                    entry.update(status="failed", error=str(exc))
                except FolderError:
                    raise
                except Exception:
                    # Never the exception text: it can quote the member's bytes.
                    entry.update(status="failed", error="processing failed")
                else:
                    # Outside the member's own error handling on purpose: a sink
                    # that fails part-way has a broken stream, so nothing after
                    # it can be written correctly and the whole run must stop.
                    sink.add_file(arcname, produced)
                    used.add(arcname)
                    entry.update(status="ok", output=arcname)
            finally:
                cleanup_member()
            results.append(entry)

        ok = sum(1 for r in results if r["status"] == "ok")
        manifest = {
            "mode": MODE_PER_FILE,
            "application": _APPLICATION_BY_FORMAT.get(fmt, fmt),
            "format": fmt,
            "paths_deidentified": False,
            "paths_note": "Paths are the uploader's own file and folder names "
                          "and have not been de-identified.",
            "files_total": len(results),
            "files_succeeded": ok,
            "files_failed": len(results) - ok,
            "files": results,
        }
        if ok == 0:
            raise FolderError(
                f"none of the {len(results)} file(s) in the folder could be de-identified"
            )
        closed = sink.close(MANIFEST_NAME, json.dumps(manifest, indent=2).encode("utf-8"))
    except BaseException:
        sink.abort()
        raise
    return manifest, closed


class SinkFailure(Exception):
    """The output sink itself failed (upload, encryption, size). Fatal to the
    whole run: unlike one member's failure, nothing after it can be written."""


class LocalTarSink:
    """The sink protocol over a plain tar file. Used where the output does not
    leave through the streaming container (tests, local tooling)."""

    def __init__(self, path: str):
        self.path = path
        self._tar = tarfile.open(path, "w:", format=tarfile.PAX_FORMAT)
        os.chmod(path, 0o600)
        self._now = int(_dt.datetime.now(tz=_dt.timezone.utc).timestamp())

    def _info(self, arcname, size):
        info = tarfile.TarInfo(arcname)
        info.size, info.mode, info.mtime = size, 0o644, self._now
        return info

    def add_file(self, arcname: str, path: str) -> None:
        with open(path, "rb") as fh:
            self._tar.addfile(self._info(arcname, os.fstat(fh.fileno()).st_size), fh)

    def close(self, manifest_name: str, manifest: bytes) -> str:
        self._tar.addfile(self._info(manifest_name, len(manifest)), io.BytesIO(manifest))
        self._tar.close()
        return self.path

    def abort(self) -> None:
        try:
            self._tar.close()
        finally:
            _unlink(self.path)


def folder_status(mode: str, files_total: int, **counts) -> dict:
    """The `outputs.folder` block of status.json."""
    block = {"mode": mode, "files_total": files_total}
    if mode == MODE_JOINT:
        block["records_in"] = counts.get("records_in", 0)
    else:
        block["files_succeeded"] = counts.get("files_succeeded", 0)
        block["files_failed"] = counts.get("files_failed", 0)
    return block


def _unlink(path: str) -> None:
    try:
        os.unlink(path)
    except FileNotFoundError:
        pass
